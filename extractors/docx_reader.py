from docx import Document

def extract_docx_text(file_path):
    """
    Extract text from DOCX paragraphs and tables.
    """

    doc = Document(file_path)

    text = ""

    # ---------- Paragraphs ----------
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    # ---------- Tables ----------
    for table in doc.tables:

        text += "\n--- Table ---\n"

        for row in table.rows:

            row_data = []

            for cell in row.cells:
                row_data.append(cell.text.strip())

            text += " | ".join(row_data) + "\n"

    return text